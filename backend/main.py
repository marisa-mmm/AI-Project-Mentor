import io
import os
import re
import json
import asyncio
import time
from datetime import datetime
from pathlib import Path
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from dotenv import load_dotenv
from groq import AsyncGroq, Groq

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfgen import canvas
from docx import Document

from backend import database

# Load .env from project root
BASE_DIR = Path(__file__).resolve().parent.parent
env_path = BASE_DIR / ".env"
load_dotenv(dotenv_path=env_path)

app = FastAPI(title="AI Project Mentor API (Groq-Powered)")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

groq_key = os.getenv("GROQ_API_KEY", "").strip().strip('"').strip("'")

# Preferred Groq model with instant fallback chain
GROQ_MODELS = [
    "llama-3.3-70b-versatile",
    "llama-3.1-8b-instant"
]
ACTIVE_MODEL = GROQ_MODELS[0]

FALLBACK_GROQ_MODELS = [
    "qwen/qwen3.8-27b",
    "groq/compound",
    "openai/gpt-oss-120b",
    "mixtral-8x7b-32768"
]

@app.on_event("startup")
def sanitize_cached_error_reports():
    err_pattern = {"$regex": "model_decommissioned|Technical Error|Error code: 400|decommissioned|gemma2", "$options": "i"}
    
    # Reset full thesis reports that cached the error
    if database.blueprints_col is not None:
        res = database.blueprints_col.update_many(
            {"full_thesis_report": err_pattern},
            {"$set": {"full_thesis_report": ""}}
        )
        print(f"Startup DB cleanup: reset {res.modified_count} stale/errored thesis reports.")
    
    # Remove any cached verifications containing the error string
    if database.db is not None and "report_verifications" in database.db.list_collection_names():
        del_res = database.db["report_verifications"].delete_many(
            {"verification": err_pattern}
        )
        print(f"Startup DB cleanup: removed {del_res.deleted_count} errored verifications.")

if groq_key:
    async_client = AsyncGroq(api_key=groq_key)
    sync_client = Groq(api_key=groq_key)
    print(f"Groq configured successfully using key ending in: ...{groq_key[-4:]}")
else:
    print(f"ERROR: GROQ_API_KEY not found in {env_path}")
    async_client = None
    sync_client = None

def clean_output(text: str) -> str:
    if not text:
        return ""
    cleaned = re.sub(r'^[ \t]*#{1,6}[ \t]*', '', text, flags=re.MULTILINE)
    return cleaned.strip()

# HIGH-SPEED ASYNC GROQ CALL WITH AUTOMATIC FALLBACK
async def query_groq_async(prompt: str, max_tokens: int = 800, fallback: str = "Report compiled successfully.") -> str:
    if not async_client:
        return fallback

    models_to_try = GROQ_MODELS + [m for m in FALLBACK_GROQ_MODELS if m not in GROQ_MODELS]
    last_err = "No compatible Groq model available."
    for model_name in models_to_try:
        try:
            response = await async_client.chat.completions.create(
                model=model_name,
                messages=[
                    {"role": "system", "content": "You are a senior academic project evaluator. Write concise, highly technical engineering project reports without markdown hashtags. Do NOT wrap package names or technologies in single backticks. Write plain text."},
                    {"role": "user", "content": prompt}
                ],
                temperature=0.3,
                max_tokens=max_tokens,
            )
            text = response.choices[0].message.content or fallback
            return clean_output(text)
        except Exception as e:
            last_err = str(e)
            if any(err_flag in last_err.lower() for err_flag in [
                "model_decommissioned", 
                "model_not_found", 
                "deprecated", 
                "invalid_request_error",
                "decommissioned",
                "400",
                "404"
            ]):
                print(f"Skipping unsupported model {model_name}: {last_err}")
                continue
            print(f"Groq async query error ({model_name}): {e}")
            continue

    return f"{fallback}\n\nTechnical Error: {last_err}"

# SYNC GROQ CALL HELPER
def query_groq_sync(prompt: str, max_tokens: int = 2500, fallback: str = "Report compiled successfully.") -> str:
    if not sync_client:
        return fallback

    models_to_try = GROQ_MODELS + [m for m in FALLBACK_GROQ_MODELS if m not in GROQ_MODELS]
    last_err = "No compatible Groq model available."
    for model_name in models_to_try:
        try:
            response = sync_client.chat.completions.create(
                model=model_name,
                messages=[
                    {"role": "system", "content": "You are a university academic project evaluator. Write clean, technical academic reports without markdown hashtags. Do NOT wrap package names or technologies in single backticks. Write plain text."},
                    {"role": "user", "content": prompt}
                ],
                temperature=0.3,
                max_tokens=max_tokens,
            )
            text = response.choices[0].message.content or fallback
            return clean_output(text)
        except Exception as e:
            last_err = str(e)
            if any(err_flag in last_err.lower() for err_flag in [
                "model_decommissioned", 
                "model_not_found", 
                "deprecated", 
                "invalid_request_error",
                "decommissioned",
                "400",
                "404"
            ]):
                print(f"Skipping unsupported model {model_name}: {last_err}")
                continue
            print(f"Groq sync query error ({model_name}): {e}")
            continue

    return f"{fallback}\n\nTechnical Error: {last_err}"

# DOMAIN DETECTOR
async def detect_project_domain_async(name: str, problem: str) -> str:
    prompt = f"""
Identify the engineering domain (2-3 words) for this project:
Project: {name}
Problem: {problem[:150]}
Reply ONLY with the domain name (e.g., Computer Vision, FinTech, EdTech, Healthcare AI, IoT Systems).
"""
    res = await query_groq_async(prompt, max_tokens=30, fallback="Applied Software Engineering")
    clean = res.replace('"', '').replace("'", "").replace(".", "").strip()
    return clean if clean else "Applied Software Engineering"

# AGENTS
async def run_idea_agent_async(name: str, problem: str, domain: str) -> str:
    prompt = f"""
You are an Academic Project Evaluator. Explain everything in simple, everyday language, but provide thorough, practical detail.
Do NOT use markdown hashtags (#). Use only the exact bold headings below:

Project Name: {name}
Problem: {problem}
Domain: {domain}

**1. Is this a good academic project?**
Provide a detailed explanation of whether this is a strong university project, assessing its technical depth, semester feasibility, and academic merit in clear terms.

**2. What problem does it solve?**
Explain the real-world difficulty or inefficiency this project fixes, and detail why existing or traditional methods are not good enough.

**3. Who will use it?**
List and detail the primary end-users, student groups, and administrative stakeholders who will interact with and benefit from this application.

**4. What can make the project innovative?**
Describe at least 3 distinct, creative ideas or unique features that will make this project stand out from ordinary tutorial projects.
"""
    return await query_groq_async(prompt, max_tokens=1500)

async def run_scope_agent_async(name: str, problem: str) -> str:
    prompt = f"""
You are an AI Project Scope Agent. Define the boundaries of this project using simple, clear, and detailed explanations.
Do NOT use markdown hashtags (#). Use only the exact bold headings below:

Project Name: {name}
Problem: {problem}

**1. Project Objective**
State the clear, measurable primary goals that the project intends to achieve.

**2. 5 Important Features**
Detail the 5 essential core features:
* **Feature 1**: Description, functionality, and how users interact with it.
* **Feature 2**: Description, functionality, and how users interact with it.
* **Feature 3**: Description, functionality, and how users interact with it.
* **Feature 4**: Description, functionality, and how users interact with it.
* **Feature 5**: Description, functionality, and how users interact with it.

**3. Project Scope**
Detail the technical workflows, modules, data pipelines, and user flows that are officially included within the system.

**4. What Should Not Be Included**
State at least 3 features or complex additions that will purposely NOT be built to prevent scope creep and protect the project schedule.

**5. Expected Output**
List the runnable software modules, user interfaces, database collections, and documentation deliverables that will be handed over.
"""
    return await query_groq_async(prompt, max_tokens=1500)

async def run_technology_agent_async(name: str, domain: str, problem: str) -> str:
    prompt = f"""
You are a Technology Selection Agent. Recommend modern, beginner-friendly yet industry-standard technologies for this project.
Explain your recommendations in simple words with thorough detail.
Do NOT use markdown hashtags (#).

Student Project: {name}
Domain: {domain}
Problem: {problem}

Format your response strictly following this structure:
* **Programming Language**: [Recommended Language] - **Why it is useful**: [Detailed explanation in simple terms explaining suitability, performance, and ecosystem]
* **AI/ML Technology**: [Recommended AI Model / Tool / Library] - **Why it is useful**: [Detailed explanation in simple terms explaining accuracy, ease of integration, and speed]
* **Database**: [Recommended Database] - **Why it is useful**: [Detailed explanation in simple terms explaining data storage, schema flexibility, and ease of querying]
* **Backend**: [Recommended Backend Framework] - **Why it is useful**: [Detailed explanation in simple terms explaining routing speed, stability, and simplicity]
* **Frontend**: [Recommended Frontend Framework] - **Why it is useful**: [Detailed explanation in simple terms explaining UI responsiveness, component setup, and styling]
* **Deployment Platform**: [Recommended Hosting Platform] - **Why it is useful**: [Detailed explanation in simple terms explaining zero-cost student tiers, deployment ease, and reliability]
"""
    return await query_groq_async(prompt, max_tokens=1500)

async def run_planning_agent_async(name: str, duration_months: int, problem: str) -> str:
    prompt = f"""
You are an Academic Project Planning Agent. Create a clear project roadmap across {duration_months} months using simple language and structured detail.
Do NOT use markdown hashtags (#).

Project Name: {name}
Duration: {duration_months} Months
Problem: {problem}

Divide the roadmap across Month 1 through Month {duration_months}. Format every month exactly as follows:

**Month [N]: [Phase Name]**
* **Task**: Detailed development tasks covering UI, backend services, or AI integration.
* **Expected Output**: Tangible, testable deliverables completed by the end of this month.
* **Testing & Verification**: Exact checks and bug tests performed during this phase.

Ensure the final month explicitly includes end-to-end integration testing and full academic documentation preparation.
"""
    return await query_groq_async(prompt, max_tokens=1800)

async def run_risk_agent_async(name: str, domain: str, duration_months: int) -> str:
    prompt = f"""
Identify technical risks and mitigations for: {name} ({duration_months} Months)
**1. Technical Risks & Failure Points**
* **Risk 1 (API & Service Limits)**: Failure description.
* **Risk 2 (Data Consistency & Storage)**: Failure description.
* **Risk 3 (Model Latency & Hallucination)**: Failure description.

**2. Execution & Timeline Risks**
* **Risk 4 (Scope Creep)**: Deadline risk.
* **Risk 5 (Integration Bottlenecks)**: Development roadblock.

**3. Actionable Mitigation Protocols**
* **Mitigation 1**: Service limits solution.
* **Mitigation 2**: Data backup/recovery protocol.
* **Mitigation 3**: AI latency/fallback guardrail.
* **Mitigation 4**: Timeline recovery strategy.
"""
    return await query_groq_async(prompt, max_tokens=700)

async def run_thesis_outline_agent_async(name: str) -> str:
    prompt = f"""
Provide ONLY the 14-Section University Thesis Outline checklist for '{name}'.
Format strictly as 14 numbered points with 1 sentence each:
**1. Abstract**: Short summary of {name}.
**2. Introduction**: Why students/users need this project.
**3. Problem Statement**: Details on user struggles and bottlenecks.
**4. Objectives**: What the system sets out to do.
**5. Literature Review**: Looking at existing tools and research gaps.
**6. Methodology**: How we plan to build the system.
**7. System Architecture**: Component diagrams and layered code layout.
**8. Technologies**: Programming languages, frameworks, and databases used.
**9. Implementation**: Writing the code, APIs, and key logic.
**10. Testing & Validation**: Checking for bugs and performance metrics.
**11. Results & Discussions**: Evaluation of results and project impact.
**12. Future Scope**: Adding advanced features and upgrades later.
**13. Conclusion**: Final thoughts on the engineering contribution.
**14. References**: List all books, IEEE citations, documentation, and websites strictly as sub-bullet points under section 14 (e.g., * Reference Title).
CRITICAL CONSTRAINT: Do NOT number references as 15, 16, 17... The final section must remain numbered 14.
"""
    return await query_groq_async(prompt, max_tokens=550)

def run_full_thesis_agent(name: str, domain: str, problem: str = "") -> str:
    prompt = f"""
You are a Senior Academic Thesis Director.
Write a comprehensive, formal 14-Section academic thesis manuscript for:
Project Title: {name}
Domain: {domain}
Problem: {problem}

You MUST write out ALL 14 numbered sections in full detail:
**1. Abstract**
**2. Introduction**
**3. Problem Statement**
**4. Objectives**
**5. Literature Review**
**6. Methodology**
**7. System Architecture**
**8. Technologies**
**9. Implementation**
**10. Testing & Validation**
**11. Results & Discussions**
**12. Future Scope**
**13. Conclusion**
**14. References**
List all books, IEEE citations, documentation, and websites strictly as sub-bullet points under section 14 (e.g., * Reference Title).
CRITICAL CONSTRAINT: Do NOT number references as 15, 16, 17... The final section must remain numbered 14.

Do NOT use markdown hashtags (#). Provide thorough academic paragraphs for each section.
"""
    return query_groq_sync(prompt, max_tokens=4000)

# CANVAS
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        if self._pageNumber == 1:
            return

        self.saveState()
        self.setFont("Helvetica-Bold", 8)
        self.setFillColor(colors.HexColor("#64748B"))

        self.drawString(54, 750, "PROJECT BLUEPRINT & SPECIFICATION")
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(54, 742, 558, 742)

        self.line(54, 50, 558, 50)
        self.setFont("Helvetica", 8)
        self.drawString(54, 38, "Academic Project Mentor — Official Specification Document")
        self.drawRightString(558, 38, f"Page {self._pageNumber} of {page_count}")
        self.restoreState()

# SCHEMAS
class RegisterSchema(BaseModel):
    username: str
    email: str
    password: str
    role: str = "Student"

class LoginSchema(BaseModel):
    email: str
    password: str

class BlueprintCreateRequest(BaseModel):
    name: str
    domain: str = ""
    duration_months: int = 3
    problem_statement: str
    user_email: str

class MentorChatRequest(BaseModel):
    project_name: str
    context: str
    query: str

class FacultyReviewSchema(BaseModel):
    project_name: str
    status: str
    comments: str

class DynamicQuestionsRequest(BaseModel):
    name: str
    idea: str
    level: str

class FacultySummaryRequest(BaseModel):
    name: str
    problem: str
    idea_eval: str
    scope_def: str

class StudentProgressUpdate(BaseModel):
    project_name: str
    user_email: str
    completion_percentage: int  # 0 to 100
    current_phase: str          # e.g., "UI Design", "Backend APIs", "Testing", "Documentation"
    notes: str = ""

class VivaGradingSchema(BaseModel):
    project_name: str
    student_email: str
    faculty_email: str
    architecture_score: int    # / 10
    code_execution_score: int  # / 20
    presentation_score: int    # / 10
    viva_qa_score: int         # / 10
    total_score: int           # / 50
    verdict: str               # "Excellent", "Satisfactory", "Re-examination"
    examiner_remarks: str

class AnnouncementSchema(BaseModel):
    title: str
    content: str
    tag: str = "General"  # e.g., "Deadline", "Viva Notice", "Report Guidelines"
    faculty_name: str

class ReportVerificationRequest(BaseModel):
    project_name: str
    manuscript_text: str

class ThesisCommentRequest(BaseModel):
    project_name: str
    faculty_email: str
    section: str  # e.g., "Abstract", "Methodology", "General"
    comment: str

class MentorAssignRequest(BaseModel):
    project_name: str
    guide_name: str
    guide_email: str

class BiWeeklyLogRequest(BaseModel):
    project_name: str
    student_email: str
    log_summary: str
    blockers: str = ""
    hours_spent: int

class GuideApprovalRequest(BaseModel):
    project_name: str
    log_index: int
    status: str  # "Approved" or "Needs Clarification"
    guide_notes: str

class PitfallRequest(BaseModel):
    project_name: str
    domain: str
    problem: str

# AUTH
@app.post("/api/auth/register")
@app.post("/register")
async def register_user(req: RegisterSchema):
    email_clean = req.email.strip().lower()
    existing = database.users_col.find_one({"email": email_clean})
    if existing:
        raise HTTPException(status_code=400, detail="An account with this email already exists.")

    new_user = {
        "username": req.username.strip(),
        "email": email_clean,
        "password": req.password,
        "role": req.role
    }
    database.users_col.insert_one(new_user)

    return {
        "success": True,
        "user": {
            "username": new_user["username"],
            "email": new_user["email"],
            "role": new_user["role"]
        }
    }

@app.post("/api/auth/login")
@app.post("/login")
async def login_user(req: LoginSchema):
    email_clean = req.email.strip().lower()
    user = database.users_col.find_one({"email": email_clean})
    if not user or user.get("password") != req.password:
        raise HTTPException(status_code=401, detail="Invalid credentials.")

    return {
        "success": True,
        "user": {
            "username": user.get("username", "Student"),
            "email": user.get("email"),
            "role": user.get("role", "Student")
        }
    }

# FAST BLUEPRINT GENERATION
@app.post("/api/generate-blueprint")
async def generate_blueprint(req: BlueprintCreateRequest):
    try:
        p_name = req.name.strip() or f"{req.problem_statement[:25]} Platform"
        p_months = int(req.duration_months) if req.duration_months else 3
        p_prob = req.problem_statement.strip()

        raw_domain = req.domain.strip() if req.domain else ""
        if not raw_domain or raw_domain.lower() in ["auto-detect", "none", "string"]:
            p_domain = await detect_project_domain_async(p_name, p_prob)
        else:
            p_domain = raw_domain

        idea_task = run_idea_agent_async(p_name, p_prob, p_domain)
        scope_task = run_scope_agent_async(p_name, p_prob)
        tech_task = run_technology_agent_async(p_name, p_domain, p_prob)
        plan_task = run_planning_agent_async(p_name, p_months, p_prob)
        risk_task = run_risk_agent_async(p_name, p_domain, p_months)
        thesis_task = run_thesis_outline_agent_async(p_name)

        idea_eval, scope_def, tech_stack, planning, risk_eval, thesis_outline = await asyncio.gather(
            idea_task, scope_task, tech_task, plan_task, risk_task, thesis_task
        )

        blueprint = {
            "user_email": req.user_email.strip().lower(),
            "name": p_name,
            "approval_status": "Pending Review",
            "project_details": {
                "name": p_name,
                "domain": p_domain,
                "duration_months": p_months,
                "problem_statement": p_prob,
                "preferred_tech": "AI Recommended Stack"
            },
            "idea_evaluation": idea_eval,
            "scope_definition": scope_def,
            "technology_stack": tech_stack,
            "time_planning": planning,
            "risk_assessment": risk_eval,
            "thesis_format": thesis_outline,
            "full_thesis_report": "",
            "timeline_milestones": planning,
            "documentation_plan": thesis_outline,
            "faculty_feedback": ""
        }

        database.save_blueprint(blueprint)
        blueprint.pop("_id", None)
        return blueprint
    except Exception as e:
        print(f"Generation error: {e}")
        raise HTTPException(status_code=500, detail=str(e))

# FACULTY AI EXECUTIVE SUMMARY
@app.post("/api/faculty/summarize")
async def summarize_blueprint_for_faculty(req: FacultySummaryRequest):
    prompt = f"""
You are an Academic Review Assistant. Faculty professors do not have time to read long blueprints.
Summarize this project blueprint in crisp, professional points:
Project Name: {req.name}
Problem: {req.problem[:300]}
Idea Feasibility: {req.idea_eval[:400]}
Scope: {req.scope_def[:400]}

Format strictly as these 4 bullet points (max 2 clear sentences each):
* **Project Core**: What the system builds and its exact functionality.
* **Academic Feasibility**: Practicality for semester completion and engineering depth.
* **Key Innovation**: The primary standout technical feature.
* **Faculty Recommendation**: Suggested verdict (Approve or Revise) with one clear justification.
"""
    summary = await query_groq_async(prompt, max_tokens=350)
    return {"summary": summary}

# FULL THESIS GENERATION
@app.get("/api/project/full-thesis")
def get_or_generate_full_thesis(project_name: str):
    query = {"project_details.name": {"$regex": f"^{re.escape(project_name.strip())}$", "$options": "i"}}
    blueprint = database.blueprints_col.find_one(query)
    if not blueprint:
        raise HTTPException(status_code=404, detail="Project not found.")

    existing_report = blueprint.get("full_thesis_report", "")
    if existing_report and not any(k in str(existing_report) for k in ["model_decommissioned", "Technical Error", "Error code: 400", "decommissioned"]):
        return {"report": existing_report}

    det = blueprint.get("project_details", {})
    report = run_full_thesis_agent(
        det.get("name", project_name),
        det.get("domain", "Computer Engineering"),
        det.get("problem_statement", "")
    )

    if not any(k in str(report) for k in ["model_decommissioned", "Technical Error", "Error code: 400", "decommissioned"]):
        database.blueprints_col.update_one(query, {"$set": {"full_thesis_report": report}})
    return {"report": report}

@app.get("/api/user/history")
async def get_user_history(email: str):
    projects = list(database.blueprints_col.find({"user_email": email.strip().lower()}, {"_id": 0}))
    return projects

@app.delete("/api/project/{project_name}")
async def delete_project(project_name: str, email: str):
    query = {
        "user_email": email.strip().lower(),
        "project_details.name": {"$regex": f"^{re.escape(project_name.strip())}$", "$options": "i"}
    }
    result = database.blueprints_col.delete_one(query)
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Project not found.")
    return {"success": True, "message": "Project deleted successfully."}

@app.post("/api/project/generate-questions")
async def generate_project_questions(req: DynamicQuestionsRequest):
    prompt = f"""
Generate exactly 3 relevant Multiple Choice Questions for an engineering student project based on:
Project Name: {req.name}
Project Idea: {req.idea}
Target Skill Level: {req.level} (Beginner = simple/foundational, Intermediate = practical choices, Advanced = moderate architecture).

Format response STRICTLY as valid JSON matching this structure without Markdown wraps:
{{
  "questions": [
    {{
      "id": "q1",
      "question": "Clear question here?",
      "options": ["Option A", "Option B", "Option C"]
    }},
    {{
      "id": "q2",
      "question": "Clear question here?",
      "options": ["Option A", "Option B", "Option C"]
    }},
    {{
      "id": "q3",
      "question": "Clear question here?",
      "options": ["Option A", "Option B", "Option C"]
    }}
  ]
}}
"""
    raw_json = await query_groq_async(prompt, max_tokens=600)
    try:
        clean_json = re.sub(r'^```json\s*|\s*```$', '', raw_json.strip(), flags=re.MULTILINE)
        return json.loads(clean_json)
    except Exception:
        if req.level == "Intermediate":
            return {
                "questions": [
                    {"id": "q_users", "question": "Which specific user roles need dedicated access in your system?", "options": ["Two distinct roles: Standard User and Administrator / Reviewer", "Multi-tier roles: Students, Teachers/Mentors, and Super-Admins", "Public visitors with optional authenticated member privileges"]},
                    {"id": "q_problem", "question": "What existing bottleneck does your project address most effectively?", "options": ["Scattered data across chat apps, spreadsheets, and emails", "High latency and lack of real-time visibility in daily tracking", "Lack of automated checks, reporting, and structured validation"]},
                    {"id": "q_features", "question": "What primary interactive capability will your platform deliver?", "options": ["Automated status workflows with live analytical dashboards", "Role-based approvals, feedback forms, and file verification", "Intelligent search engine with history logs and data export"]}
                ]
            }
        elif req.level == "Advanced":
            return {
                "questions": [
                    {"id": "q_problem", "question": "What core engineering challenge does your problem statement target?", "options": ["Managing high-frequency data updates with zero inconsistency", "Integrating automated decision logic and recommendation pipelines", "Ensuring secure cross-platform data synchronization and export"]},
                    {"id": "q_features", "question": "Which advanced capability should be prioritized in the MVP?", "options": ["Real-time data synchronization with instant status reflection", "Automated insights, summary generation, and metric reporting", "Role-guarded audit trails, activity logging, and PDF/Excel reports"]},
                    {"id": "q_innovation", "question": "What key innovation will make this capstone stand out during viva defense?", "options": ["Smart recommendation / AI assistant embedded in the workflow", "Clean modular service architecture with automated health checks", "End-to-end responsive design with offline caching support"]}
                ]
            }
        return {
            "questions": [
                {"id": "q_users", "question": "Who are the primary target users for this project?", "options": ["College students and campus faculty", "General public / everyday consumers", "Small businesses and administrative staff"]},
                {"id": "q_problem", "question": "What is the main problem your application focuses on solving?", "options": ["Automating manual paperwork and repetitive tasks", "Providing quick access to resources, notes, or records", "Improving communication, alerts, and tracking between people"]},
                {"id": "q_features", "question": "Which core feature is essential to launch your project?", "options": ["User dashboard with search, filters, and detail views", "Instant alert notifications, status updates, and reminders", "Secure file, image, or document submission and storage"]}
            ]
        }

@app.get("/api/faculty/blueprints")
async def get_all_blueprints():
    projects = list(database.blueprints_col.find({}, {"_id": 0}))
    return projects

@app.post("/api/faculty/review")
async def submit_faculty_review(rev: FacultyReviewSchema):
    query = {"project_details.name": {"$regex": f"^{re.escape(rev.project_name.strip())}$", "$options": "i"}}
    update = {
        "$set": {
            "approval_status": rev.status,
            "faculty_feedback": rev.comments
        }
    }
    result = database.blueprints_col.update_one(query, update)
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Project not found.")
    return {"success": True, "message": "Evaluation recorded."}

@app.post("/api/student/update-progress")
async def update_student_progress(req: StudentProgressUpdate):
    query = {
        "user_email": req.user_email.strip().lower(),
        "project_details.name": {"$regex": f"^{re.escape(req.project_name.strip())}$", "$options": "i"}
    }
    update = {
        "$set": {
            "progress.completion_percentage": req.completion_percentage,
            "progress.current_phase": req.current_phase,
            "progress.notes": req.notes,
            "progress.last_updated": "Today"
        }
    }
    if database.blueprints_col is not None:
        database.blueprints_col.update_one(query, update)
    return {"success": True, "message": "Progress updated successfully."}

@app.post("/api/faculty/viva-grading")
async def submit_viva_grading(req: VivaGradingSchema):
    if database.db is not None:
        database.db["viva_scores"].update_one(
            {"project_name": req.project_name, "student_email": req.student_email},
            {"$set": req.dict()},
            upsert=True
        )
    return {"success": True, "message": "Viva grading saved."}

@app.get("/api/faculty/viva-scores")
async def get_viva_scores():
    if database.db is not None:
        return list(database.db["viva_scores"].find({}, {"_id": 0}))
    return []

@app.post("/api/faculty/announcements")
async def create_announcement(req: AnnouncementSchema):
    doc = req.dict()
    doc["created_at"] = datetime.now().strftime("%b %d, %Y • %I:%M %p")
    doc["timestamp"] = int(time.time() * 1000)
    if database.db is not None:
        database.db["announcements"].insert_one(doc)
        doc.pop("_id", None)
    return {"success": True, "message": "Notice posted."}

@app.get("/api/announcements")
async def get_announcements():
    if database.db is not None:
        items = list(database.db["announcements"].find({}, {"_id": 0}))
        items.sort(key=lambda x: x.get("timestamp", 0), reverse=True)
        return items
    return []

@app.post("/api/faculty/thesis-comment")
async def add_thesis_comment(req: ThesisCommentRequest):
    query = {
        "$or": [
            {"project_details.name": {"$regex": f"^{re.escape(req.project_name.strip())}$", "$options": "i"}},
            {"name": {"$regex": f"^{re.escape(req.project_name.strip())}$", "$options": "i"}}
        ]
    }
    comment_data = {
        "faculty_email": req.faculty_email,
        "section": req.section,
        "comment": req.comment,
        "created_at": datetime.now().strftime("%b %d, %Y • %I:%M %p")
    }
    if database.blueprints_col is not None:
        database.blueprints_col.update_one(
            query,
            {"$push": {"thesis_comments": comment_data}}
        )
    return {"success": True, "message": "Feedback saved."}

@app.post("/api/faculty/verify-report")
async def verify_thesis_report(req: ReportVerificationRequest):
    prompt = f"""
You are an Academic Examination Board Verifier. Check this university student report text:
Project: {req.project_name}
Text: {req.manuscript_text[:3500]}

Evaluate and return strictly 4 points:
* **Structural Completeness**: Are major sections (Abstract, Problem, Tech, Testing) covered?
* **Originality & Novelty Check**: Assess if content looks genuine or boilerplate.
* **Citation & References Quality**: Evaluate academic referencing.
* **Verification Score**: Provide an estimated compliance score (e.g., 88/100) and recommendation (Accepted / Resubmit).
"""
    result = await query_groq_async(prompt, max_tokens=600)
    return {"verification": result}

@app.post("/api/mentor/chat")
async def mentor_chat(req: MentorChatRequest):
    prompt = f"""
You are the Lead Academic Project Mentor.
Project: {req.project_name}
Context: {req.context}
Student Question: {req.query}

Provide a direct, practical response with clean formatting and bold highlights.
Do NOT use hashtags (# or ##).
"""
    reply = query_groq_sync(prompt, max_tokens=1000)
    return {"reply": reply}

@app.post("/api/mentor/assign")
async def assign_mentor(req: MentorAssignRequest):
    if database.blueprints_col is None:
        raise HTTPException(status_code=500, detail="Database not connected.")
    query = {"project_details.name": {"$regex": f"^{re.escape(req.project_name.strip())}$", "$options": "i"}}
    result = database.blueprints_col.update_one(query, {"$set": {
        "assigned_guide": {
            "guide_name": req.guide_name,
            "guide_email": req.guide_email.strip().lower(),
            "status": "Assigned"
        }
    }})
    if result.matched_count == 0:
        database.blueprints_col.update_one(
            {"name": {"$regex": f"^{re.escape(req.project_name.strip())}$", "$options": "i"}},
            {"$set": {
                "assigned_guide": {
                    "guide_name": req.guide_name,
                    "guide_email": req.guide_email.strip().lower(),
                    "status": "Assigned"
                }
            }}
        )
    return {"success": True, "message": "Guide assigned."}

@app.post("/api/student/submit-log")
async def submit_biweekly_log(req: BiWeeklyLogRequest):
    if database.blueprints_col is None:
        raise HTTPException(status_code=500, detail="Database not connected.")
    query = {"project_details.name": {"$regex": f"^{re.escape(req.project_name.strip())}$", "$options": "i"}}
    log_entry = {
        "summary": req.log_summary,
        "blockers": req.blockers,
        "hours_spent": req.hours_spent,
        "status": "Pending Guide Approval",
        "guide_notes": "",
        "timestamp": datetime.now().strftime("%b %d, %Y • %I:%M %p")
    }
    result = database.blueprints_col.update_one(query, {"$push": {"mentor_logs": log_entry}})
    if result.matched_count == 0:
        database.blueprints_col.update_one(
            {"name": {"$regex": f"^{re.escape(req.project_name.strip())}$", "$options": "i"}},
            {"$push": {"mentor_logs": log_entry}}
        )
    return {"success": True, "message": "Log submitted to mentor."}

@app.post("/api/mentor/verify-log")
async def verify_biweekly_log(req: GuideApprovalRequest):
    if database.blueprints_col is None:
        raise HTTPException(status_code=500, detail="Database not connected.")
    query = {"project_details.name": {"$regex": f"^{re.escape(req.project_name.strip())}$", "$options": "i"}}
    result = database.blueprints_col.update_one(
        query,
        {"$set": {
            f"mentor_logs.{req.log_index}.status": req.status,
            f"mentor_logs.{req.log_index}.guide_notes": req.guide_notes
        }}
    )
    if result.matched_count == 0:
        database.blueprints_col.update_one(
            {"name": {"$regex": f"^{re.escape(req.project_name.strip())}$", "$options": "i"}},
            {"$set": {
                f"mentor_logs.{req.log_index}.status": req.status,
                f"mentor_logs.{req.log_index}.guide_notes": req.guide_notes
            }}
        )
    return {"success": True, "message": "Log evaluation recorded."}

@app.post("/api/projects/pitfalls")
async def get_project_pitfalls(req: PitfallRequest):
    prompt = f"""
You are a Senior Academic Project Evaluator.
Identify the 2 most common technical mistakes or architectural pitfalls students make in '{req.domain}' projects like '{req.project_name}'.
Explain why previous batches faced rejection or heavy revisions on this concept, and what a student must do to pass review cleanly.
Keep it strictly to 2 short, impactful bullet points. Plain text, no hashtags.
"""
    result = await query_groq_async(prompt, max_tokens=250)
    return {"pitfalls": result}

# EXPORT PDF
@app.get("/api/export/pdf")
async def export_blueprint_pdf(project_name: str):
    query = {"project_details.name": {"$regex": f"^{re.escape(project_name.strip())}$", "$options": "i"}}
    blueprint = database.blueprints_col.find_one(query, {"_id": 0})
    if not blueprint:
        raise HTTPException(status_code=404, detail=f"Blueprint for '{project_name}' not found.")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=54, rightMargin=54, topMargin=64, bottomMargin=64)
    styles = getSampleStyleSheet()

    cover_inst_style = ParagraphStyle('CoverInst', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=13, leading=17, alignment=1, textColor=colors.HexColor("#1E3A8A"), spaceAfter=12)
    cover_title_style = ParagraphStyle('CoverTitle', parent=styles['Title'], fontName='Helvetica-Bold', fontSize=24, leading=30, alignment=1, textColor=colors.HexColor("#0F172A"), spaceAfter=14)
    cover_sub_style = ParagraphStyle('CoverSub', parent=styles['Normal'], fontName='Helvetica', fontSize=11, leading=16, alignment=1, textColor=colors.HexColor("#475569"), spaceAfter=30)
    cover_meta_style = ParagraphStyle('CoverMeta', parent=styles['Normal'], fontName='Helvetica', fontSize=10, leading=16, alignment=1, textColor=colors.HexColor("#1E293B"))

    h1_style = ParagraphStyle('AcademicH1', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=13, leading=17, textColor=colors.HexColor("#0F172A"), spaceBefore=14, spaceAfter=8, keepWithNext=True)
    h2_style = ParagraphStyle('AcademicH2', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=11, leading=15, textColor=colors.HexColor("#1E40AF"), spaceBefore=8, spaceAfter=4, keepWithNext=True)
    body_style = ParagraphStyle('AcademicBody', parent=styles['Normal'], fontName='Helvetica', fontSize=9.5, leading=14.5, textColor=colors.HexColor("#1E293B"), spaceAfter=6)
    bullet_style = ParagraphStyle('AcademicBullet', parent=styles['Normal'], fontName='Helvetica', fontSize=9.5, leading=14, leftIndent=16, firstLineIndent=-10, textColor=colors.HexColor("#334155"), spaceAfter=4)

    det = blueprint.get("project_details", {})
    name = det.get('name', 'Capstone Project')
    domain = det.get('domain', 'Computer Engineering')
    duration = det.get('duration_months', 3)
    user_email = blueprint.get('user_email', 'Candidate')

    story = []
    story.append(Spacer(1, 40))
    story.append(Paragraph("PROJECT ARCHITECTURE & BLUEPRINT SPECIFICATION", cover_inst_style))
    story.append(HRFlowable(width="60%", thickness=1.5, color=colors.HexColor("#2563EB"), spaceAfter=24, spaceBefore=6))
    story.append(Paragraph(name.upper(), cover_title_style))
    story.append(Paragraph(f"Domain: <b>{domain}</b> &nbsp;|&nbsp; Target Timeline: <b>{duration} Months</b>", cover_sub_style))
    story.append(Spacer(1, 80))

    meta_text = f"""
    <b>Candidate / Submitter:</b> {user_email}<br/>
    <b>Approval Status:</b> {blueprint.get('approval_status', 'Pending Final Review')}<br/>
    <b>Project Verification:</b> Multi-Agent Architectural Committee
    """
    story.append(Paragraph(meta_text, cover_meta_style))
    story.append(Spacer(1, 130))
    story.append(Paragraph("Comprehensive Project Blueprint Generated via AI Academic Committee", ParagraphStyle('CoverFooter', fontName='Helvetica-Oblique', fontSize=9, alignment=1, textColor=colors.HexColor("#64748B"))))
    story.append(PageBreak())

    sections = [
        ("SECTION 1: IDEA FEASIBILITY & EVALUATION", blueprint.get("idea_evaluation", "")),
        ("SECTION 2: FUNCTIONAL SCOPE & BOUNDARIES", blueprint.get("scope_definition", "")),
        ("SECTION 3: ARCHITECTURAL TECHNOLOGY STACK", blueprint.get("technology_stack", "")),
        ("SECTION 4: SPRINT PLANNING & MILESTONE ROADMAP", blueprint.get("time_planning", "") or blueprint.get("timeline_milestones", "")),
        ("SECTION 5: RISK ASSESSMENT & MITIGATION", blueprint.get("risk_assessment", "")),
        ("SECTION 6: 14-SECTION UNIVERSITY THESIS FORMAT OUTLINE", blueprint.get("thesis_format", "") or blueprint.get("documentation_plan", ""))
    ]

    for section_title, content in sections:
        story.append(Paragraph(section_title, h1_style))
        story.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor("#94A3B8"), spaceAfter=10, spaceBefore=2))

        lines = str(content).split("\n")
        for line in lines:
            trimmed = line.strip()
            if not trimmed:
                story.append(Spacer(1, 4))
                continue

            # Strip non-standard unicode characters
            sanitized = re.sub(r'[\u2010\u2011\u2012\u2013\u2014\u2015\uFE58\uFE63\uFF0D]', '-', trimmed)
            sanitized = re.sub(r'[\u00A0\u2000-\u200B\u202F\u205F\u3000]', ' ', sanitized)
            sanitized = sanitized.replace('“', '"').replace('”', '"').replace('‘', "'").replace('’', "'")

            is_bullet = sanitized.startswith("* ") or sanitized.startswith("- ") or sanitized.startswith("• ")
            clean_line = re.sub(r'^[*\-•]\s+', '', sanitized)

            if '**' in clean_line:
                if clean_line.count('**') % 2 != 0:
                    parts = clean_line.split('**', 1)
                    clean_line = f"<b>{parts[0]}</b>{parts[1]}"
                else:
                    clean_line = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', clean_line)

            clean_line = clean_line.replace("**", "").replace("<", "&lt;").replace(">", "&gt;")
            clean_line = clean_line.replace("&lt;b&gt;", "<b>").replace("&lt;/b&gt;", "</b>")

            is_subheading = bool(
                re.match(r'^\d+\.\s+[A-Z]', clean_line) or 
                re.match(r'^Month\s+\d+:', clean_line) or 
                (clean_line.startswith('<b>') and clean_line.endswith('</b>') and len(clean_line) < 90)
            )

            if is_subheading and len(clean_line) < 100:
                story.append(Spacer(1, 4))
                story.append(Paragraph(clean_line, h2_style))
            elif is_bullet:
                story.append(Paragraph(f"&bull;&nbsp;&nbsp;{clean_line}", bullet_style))
            else:
                story.append(Paragraph(clean_line, body_style))

        story.append(Spacer(1, 14))

    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    clean_filename = re.sub(r'[^a-zA-Z0-9_-]', '_', name)
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={clean_filename}_Blueprint.pdf"}
    )

# EXPORT WORD
@app.get("/api/export/docx")
async def export_blueprint_docx(project_name: str):
    query = {"project_details.name": {"$regex": f"^{re.escape(project_name.strip())}$", "$options": "i"}}
    blueprint = database.blueprints_col.find_one(query, {"_id": 0})
    if not blueprint:
        raise HTTPException(status_code=404, detail=f"Blueprint for '{project_name}' not found.")

    doc = Document()
    det = blueprint.get("project_details", {})

    doc.add_heading(f"Project Blueprint: {det.get('name')}", 0)
    doc.add_paragraph(f"Domain: {det.get('domain')} | Duration: {det.get('duration_months')} Months\n")

    sections = [
        ("SECTION 1: IDEA FEASIBILITY & EVALUATION", blueprint.get("idea_evaluation", "")),
        ("SECTION 2: FUNCTIONAL SCOPE & BOUNDARIES", blueprint.get("scope_definition", "")),
        ("SECTION 3: ARCHITECTURAL TECHNOLOGY STACK", blueprint.get("technology_stack", "")),
        ("SECTION 4: SPRINT PLANNING & MILESTONE ROADMAP", blueprint.get("time_planning", "") or blueprint.get("timeline_milestones", "")),
        ("SECTION 5: RISK ASSESSMENT & MITIGATION", blueprint.get("risk_assessment", "")),
        ("SECTION 6: 14-SECTION UNIVERSITY THESIS FORMAT OUTLINE", blueprint.get("thesis_format", "") or blueprint.get("documentation_plan", ""))
    ]

    for heading, text in sections:
        doc.add_heading(heading, level=1)
        for line in str(text).split("\n"):
            clean = line.replace("**", "").replace("###", "").replace("##", "").replace("#", "").strip()
            if clean:
                doc.add_paragraph(clean)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    clean_filename = re.sub(r'[^a-zA-Z0-9_-]', '_', det.get('name', 'project'))
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={clean_filename}_Blueprint.docx"}
    )

# FULL 14-SECTION MANUSCRIPT PDF EXPORT
@app.get("/api/export/thesis-pdf")
def export_full_thesis_pdf(project_name: str):
    query = {"project_details.name": {"$regex": f"^{re.escape(project_name.strip())}$", "$options": "i"}}
    blueprint = database.blueprints_col.find_one(query)
    if not blueprint:
        raise HTTPException(status_code=404, detail=f"Project '{project_name}' not found.")

    det = blueprint.get("project_details", {})
    name = det.get('name', project_name)
    domain = det.get('domain', 'Computer Engineering')
    duration = det.get('duration_months', 3)
    user_email = blueprint.get('user_email', 'Candidate')

    report_text = blueprint.get("full_thesis_report")
    if not report_text or any(k in str(report_text) for k in ["model_decommissioned", "Technical Error", "Error code: 400", "decommissioned"]):
        report_text = run_full_thesis_agent(name, domain, det.get("problem_statement", ""))
        if not any(k in str(report_text) for k in ["model_decommissioned", "Technical Error", "Error code: 400", "decommissioned"]):
            database.blueprints_col.update_one(query, {"$set": {"full_thesis_report": report_text}})

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=54, rightMargin=54, topMargin=64, bottomMargin=64)
    styles = getSampleStyleSheet()

    cover_inst_style = ParagraphStyle('CoverInst', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=13, leading=17, alignment=1, textColor=colors.HexColor("#1E3A8A"), spaceAfter=12)
    cover_title_style = ParagraphStyle('CoverTitle', parent=styles['Title'], fontName='Helvetica-Bold', fontSize=24, leading=30, alignment=1, textColor=colors.HexColor("#0F172A"), spaceAfter=14)
    cover_sub_style = ParagraphStyle('CoverSub', parent=styles['Normal'], fontName='Helvetica', fontSize=11, leading=16, alignment=1, textColor=colors.HexColor("#475569"), spaceAfter=30)
    cover_meta_style = ParagraphStyle('CoverMeta', parent=styles['Normal'], fontName='Helvetica', fontSize=10, leading=16, alignment=1, textColor=colors.HexColor("#1E293B"))

    h1_style = ParagraphStyle('AcademicH1', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=13, leading=17, textColor=colors.HexColor("#1E3A8A"), spaceBefore=16, spaceAfter=8, keepWithNext=True)
    body_style = ParagraphStyle('AcademicBody', parent=styles['Normal'], fontName='Helvetica', fontSize=9.5, leading=14.5, textColor=colors.HexColor("#1E293B"), spaceAfter=6)
    bullet_style = ParagraphStyle('AcademicBullet', parent=styles['Normal'], fontName='Helvetica', fontSize=9.5, leading=14, leftIndent=16, firstLineIndent=-10, textColor=colors.HexColor("#334155"), spaceAfter=4)

    story = []
    story.append(Spacer(1, 40))
    story.append(Paragraph("14-SECTION UNIVERSITY THESIS REPORT", cover_inst_style))
    story.append(HRFlowable(width="60%", thickness=1.5, color=colors.HexColor("#2563EB"), spaceAfter=24, spaceBefore=6))
    story.append(Paragraph(name.upper(), cover_title_style))
    story.append(Paragraph(f"Domain: <b>{domain}</b> &nbsp;|&nbsp; Duration: <b>{duration} Months</b>", cover_sub_style))
    story.append(Spacer(1, 80))

    meta_text = f"""
    <b>Student Submitter:</b> {user_email}<br/>
    <b>Approval Status:</b> {blueprint.get('approval_status', 'Pending Final Review')}<br/>
    <b>Format Specification:</b> 14-Section University Academic Thesis Outline
    """
    story.append(Paragraph(meta_text, cover_meta_style))
    story.append(Spacer(1, 130))
    story.append(Paragraph("A Technical Thesis Report Submitted in Partial Fulfillment of Academic Requirements", ParagraphStyle('CoverFooter', fontName='Helvetica-Oblique', fontSize=9, alignment=1, textColor=colors.HexColor("#64748B"))))
    story.append(PageBreak())

    lines = str(report_text).split("\n")
    for line in lines:
        trimmed = line.strip()
        if not trimmed:
            story.append(Spacer(1, 4))
            continue

        sanitized = re.sub(r'[\u2010\u2011\u2012\u2013\u2014\u2015\uFE58\uFE63\uFF0D]', '-', trimmed)
        sanitized = re.sub(r'[\u00A0\u2000-\u200B\u202F\u205F\u3000]', ' ', sanitized)
        sanitized = sanitized.replace('“', '"').replace('”', '"').replace('‘', "'").replace('’', "'")

        if re.match(r'^\*{0,2}\d{1,2}\.\s+.*?\*{0,2}:?$', sanitized):
            clean_heading = sanitized.replace("**", "").rstrip(":").strip()
            story.append(Spacer(1, 8))
            story.append(Paragraph(clean_heading, h1_style))
            story.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor("#CBD5E1"), spaceAfter=8, spaceBefore=2))
            continue

        is_bullet = sanitized.startswith("* ") or sanitized.startswith("- ") or sanitized.startswith("• ")
        clean_line = re.sub(r'^[*\-•]\s+', '', sanitized)

        if '**' in clean_line:
            if clean_line.count('**') % 2 != 0:
                parts = clean_line.split('**', 1)
                clean_line = f"<b>{parts[0]}</b>{parts[1]}"
            else:
                clean_line = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', clean_line)

        clean_line = clean_line.replace("**", "").replace("<", "&lt;").replace(">", "&gt;")
        clean_line = clean_line.replace("&lt;b&gt;", "<b>").replace("&lt;/b&gt;", "</b>")

        if is_bullet:
            story.append(Paragraph(f"&bull;&nbsp;&nbsp;{clean_line}", bullet_style))
        else:
            story.append(Paragraph(clean_line, body_style))

    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    clean_filename = re.sub(r'[^a-zA-Z0-9_-]', '_', name)
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={clean_filename}_Full_Thesis.pdf"}
    )